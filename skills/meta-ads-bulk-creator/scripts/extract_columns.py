#!/usr/bin/env python3
"""
Extracts the structure of AdsManagerTemplate_v2.3.xltx into:
  - scripts/spec.json   (machine-readable: column order, validation enums, per-column metadata)
  - references/columns.md (human-readable reference, grouped by level)

This runs once when the bundled .xltx changes. The bulk-creator script reads
spec.json at build time -- it never round-trips the .xltx through openpyxl
(openpyxl strips data validation rules on save with a warning).

Usage:
    python scripts/extract_columns.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

SKILL_ROOT = Path(__file__).resolve().parent.parent
XLTX_PATH = SKILL_ROOT / "assets" / "AdsManagerTemplate_v2.3.xltx"
SPEC_OUT = SKILL_ROOT / "scripts" / "spec.json"
MD_OUT = SKILL_ROOT / "references" / "columns.md"

LEVEL_RULES: list[tuple[str, str]] = [
    ("Campaign", r"^(Campaign\b|Special Ad Categor|Buying Type|Tags)"),
    ("Ad Set", r"^Ad Set\b"),
    ("Ad", r"^(Ad ID|Ad Status|Ad Name)$"),
    ("Targeting", r"^(Countries|Global Regions|Excluded Global Regions|Cities|Regions|Zip|Gender|Age (Min|Max)|Education Status|College (Start|End) Year|Interested In|Relationship|Connections|Excluded Connections|Friends of Connections|Locales|Broad Category Clusters|Custom Audiences|Excluded Custom Audiences|Location Cluster IDs|Excluded Location Cluster IDs|(Large|Medium|Small) Geo Areas|Excluded (Large|Medium|Small) Geo Areas|Metro Areas|Excluded Metro Areas|Subcities|Excluded Subcities|Neighborhoods|Excluded Neighborhoods|Subneighborhoods|Excluded Subneighborhoods)$"),
    ("Placements", r"^(Publisher Platforms|Device Platforms|(Facebook|Instagram|Messenger|Oculus|Audience Network|WhatsApp) Positions)$"),
    ("Delivery", r"^(Optimization Goal|Billing Event|Bid Amount|Minimum ROAS)$"),
    ("Compliance", r"(Beneficiary|Payer|Advertiser).*"),
    ("Creative", r"^(Title|Body|Link Description|Display Link|Image Hash|Creative Type|URL Tags|Image File Name|Creative Optimization|Call to Action|Story ID|Link|Application ID|Link Object ID)$"),
    ("Carousel", r"^Product \d+"),
    ("Marketing Messages", r"^Marketing Message"),
]


def detect_level(name: str) -> str:
    for level, pattern in LEVEL_RULES:
        if re.match(pattern, name):
            return level
    return "Other"


def parse_comment(comment_text: str | None) -> dict[str, str]:
    """Cell comments follow the pattern:
        <Field name>
        <description>
        Format
        : <format>
        Required
        : <required>
    Extract description / format / required as best we can.
    """
    if not comment_text:
        return {}
    text = comment_text.strip()
    out: dict[str, str] = {}

    fmt_match = re.search(r"Format\s*:\s*(.+?)(?:\n\s*Required|\Z)", text, re.S)
    req_match = re.search(r"Required\s*:\s*(.+?)(?:\Z)", text, re.S)

    lines = text.splitlines()
    desc_lines: list[str] = []
    for i, line in enumerate(lines[1:], start=1):
        stripped = line.strip()
        if stripped in ("Format", "Required"):
            break
        if stripped.startswith(("Format", "Required")):
            break
        desc_lines.append(stripped)
    description = " ".join(l for l in desc_lines if l).strip()

    if description:
        out["description"] = description
    if fmt_match:
        out["format"] = re.sub(r"\s+", " ", fmt_match.group(1)).strip()
    if req_match:
        out["required"] = re.sub(r"\s+", " ", req_match.group(1)).strip()
    return out


def parse_dv_range(sqref: str) -> tuple[str, str]:
    """Convert e.g. 'VALIDATION!$A$2:$D$2' to ('A', 'D') row range info ignored."""
    m = re.match(r".*?\$?([A-Z]+)\$?\d+:\$?([A-Z]+)\$?\d+", sqref)
    if not m:
        return ("A", "A")
    return (m.group(1), m.group(2))


def main() -> None:
    if not XLTX_PATH.exists():
        raise SystemExit(f"Template not found at {XLTX_PATH}")

    wb = load_workbook(XLTX_PATH, data_only=False, keep_vba=False)
    main_sheet = wb["Ads Manager Template"]
    validation_sheet = wb["VALIDATION"]

    headers: list[str] = []
    for cell in main_sheet[1]:
        if cell.value is None:
            break
        headers.append(str(cell.value))

    n_cols = len(headers)

    enum_rows: dict[int, list[str]] = {}
    for row in validation_sheet.iter_rows(values_only=False):
        row_idx = row[0].row
        values = [str(c.value).strip() for c in row if c.value is not None and str(c.value).strip()]
        if values:
            enum_rows[row_idx] = values

    column_validations: dict[str, dict] = {}
    for dv in main_sheet.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        formula = dv.formula1.strip().lstrip("=")
        m = re.match(r"VALIDATION!?\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", formula)
        if not m:
            continue
        start_col, start_row, end_col, end_row = m.groups()
        if start_row != end_row:
            continue
        row_idx = int(start_row)
        values = enum_rows.get(row_idx, [])
        if not values:
            continue
        for sqref in (dv.sqref.ranges if hasattr(dv.sqref, "ranges") else [dv.sqref]):
            sqref_str = str(sqref)
            col_match = re.match(r"([A-Z]+)\d+:?([A-Z]*)?", sqref_str)
            if not col_match:
                continue
            col_letter = col_match.group(1)
            try:
                col_idx = column_index_from_string(col_letter) - 1
                if 0 <= col_idx < n_cols:
                    column_validations[headers[col_idx]] = {
                        "validation_row": row_idx,
                        "valid_values": values,
                    }
            except ValueError:
                pass

    columns: list[dict] = []
    for i, name in enumerate(headers):
        cell = main_sheet.cell(row=1, column=i + 1)
        comment = cell.comment.text if cell.comment else None
        meta = parse_comment(comment)
        col_letter = get_column_letter(i + 1)
        col_def: dict = {
            "index": i,
            "letter": col_letter,
            "name": name,
            "level": detect_level(name),
        }
        if meta.get("description"):
            col_def["description"] = meta["description"]
        if meta.get("format"):
            col_def["format"] = meta["format"]
        if meta.get("required"):
            col_def["required"] = meta["required"]
        if name in column_validations:
            col_def["valid_values"] = column_validations[name]["valid_values"]
        columns.append(col_def)

    spec = {
        "template_file": "AdsManagerTemplate_v2.3.xltx",
        "sheet_name": "Ads Manager Template",
        "n_columns": n_cols,
        "encoding": {
            "txt_encoding": "utf-16-le",
            "txt_bom": True,
            "txt_separator": "\t",
            "txt_line_ending": "\r\n",
        },
        "limits_2026": {
            "name_max_chars": 35,
            "title_warn_chars": 27,
            "title_max_chars": 40,
            "body_warn_chars": 125,
            "body_max_chars": 2200,
            "description_warn_chars": 30,
            "description_max_chars": 200,
            "import_file_max_bytes": 2_000_000,
        },
        "columns": columns,
    }

    SPEC_OUT.parent.mkdir(parents=True, exist_ok=True)
    SPEC_OUT.write_text(json.dumps(spec, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    md_lines: list[str] = [
        "# Columns reference",
        "",
        f"Generated from `assets/{XLTX_PATH.name}` by `scripts/extract_columns.py`.",
        f"Total columns: **{n_cols}** (sheet `{spec['sheet_name']}`, range A1:{get_column_letter(n_cols)}1).",
        "",
        "Re-run `python scripts/extract_columns.py` whenever the bundled template changes.",
        "",
    ]

    levels_in_order = [
        "Campaign",
        "Ad Set",
        "Targeting",
        "Placements",
        "Delivery",
        "Compliance",
        "Ad",
        "Creative",
        "Carousel",
        "Marketing Messages",
        "Other",
    ]
    grouped: dict[str, list[dict]] = {lvl: [] for lvl in levels_in_order}
    for col in columns:
        grouped.setdefault(col["level"], []).append(col)

    for lvl in levels_in_order:
        cols_in_level = grouped.get(lvl, [])
        if not cols_in_level:
            continue
        md_lines.append(f"## {lvl} ({len(cols_in_level)} columns)")
        md_lines.append("")
        md_lines.append("| Col | Name | Format | Required | Valid values |")
        md_lines.append("|-----|------|--------|----------|--------------|")
        for col in cols_in_level:
            fmt = col.get("format", "").replace("|", "\\|")
            req = col.get("required", "").replace("|", "\\|")
            vals = ", ".join(col.get("valid_values", []))
            if len(vals) > 200:
                vals = vals[:200] + " …"
            vals = vals.replace("|", "\\|")
            md_lines.append(f"| {col['letter']} | {col['name']} | {fmt} | {req} | {vals} |")
        md_lines.append("")

    MD_OUT.parent.mkdir(parents=True, exist_ok=True)
    MD_OUT.write_text("\n".join(md_lines), encoding="utf-8")

    print(f"Wrote {SPEC_OUT.relative_to(SKILL_ROOT)} ({n_cols} columns)")
    print(f"Wrote {MD_OUT.relative_to(SKILL_ROOT)}")


if __name__ == "__main__":
    main()
